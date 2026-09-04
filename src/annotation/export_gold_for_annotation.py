"""
export_gold_for_annotation.py — Export Excel des 500 Paragraphes GOLD
═══════════════════════════════════════════════════════════════════════════════════════
Phase 3a du pipeline FraFin-Reasoning.

RÔLE DANS LE PIPELINE :
    frafin_sample_rigorous_{TS}.parquet  (15 000, Phase 2)
              │
              ▼
    export_gold_for_annotation.py   ← CE SCRIPT
              │
              ▼
    frafin_gold_500_{TS}.xlsx       (vous annotez à la main)
              │
              ▼
    annotation_mistral_mass.py      (annote les ~14 500 restants
                                      + ré-annote les 200 IAA indépendamment)
              │
              ▼
    validation_iaa.py               (Kappa de Cohen humain vs Mistral sur les 200)

MÉTHODOLOGIE DE SÉLECTION DES 500 GOLD — même rigueur que la Phase 2 :
    On NE sélectionne PAS les 500 "meilleurs" paragraphes (biais de sélection).
    On applique une stratification proportionnelle sur 2 axes :
        year_bucket × csrd_quartile  (16 strates)
    avec diversification round-robin par doc_group à l'intérieur de chaque strate,
    et un plancher minimum garanti par strate.

    Le sous-ensemble de 500 reste donc représentatif de la distribution réelle
    du corpus AMF (y compris le quartile q0 "pas de CSRD"), ce qui en fait un
    Hold-out Test Set valide pour l'évaluation du modèle fine-tuné — un Gold
    Standard composé uniquement des cas "intéressants" serait inutilisable
    comme test set (sur-estimerait la performance du modèle sur des cas faciles
    ou pré-sélectionnés comme évidents).

PROTOCOLE D'ANNOTATION À L'AVEUGLE (Blind Annotation) :
    Les colonnes csrd_quartile, csrd_score_raw et stratum_id qui ont SERVI à
    sélectionner l'échantillon NE SONT PAS affichées dans la feuille d'annotation
    principale. Elles sont reléguées à une feuille interne masquée
    ("Metadonnees_Stratification"). Objectif : éviter que l'annotateur humain
    soit influencé ("ce paragraphe a été classé q3_high, je dois donc trouver
    du CSRD dedans") — un biais d'ancrage qui invaliderait le Gold Standard.

SOUS-ENSEMBLE IAA (200 / 500) :
    200 des 500 paragraphes sont marqués `iaa_subset=True` (colonne informative,
    lecture seule). Ce sont ceux qui seront ÉGALEMENT annotés indépendamment par
    Mistral Large dans annotation_mistral_mass.py, pour calculer le Kappa de
    Cohen dans validation_iaa.py. Cette information n'introduit pas de biais :
    elle ne change pas la réponse correcte, seulement quelles lignes serviront
    à la validation statistique inter-annotateurs.

Dépendances :
    pip install pandas openpyxl pyarrow

Usage :
    python export_gold_for_annotation.py
    python export_gold_for_annotation.py --input data/frafin_sample_rigorous_XXX.parquet
    python export_gold_for_annotation.py --n-gold 500 --n-iaa 200 --seed 42
"""


from __future__ import annotations

import argparse
import json
import logging
import sys
from collections import deque
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

try:
    from src.annotation.csrd_taxonomy import (
        CSRD_TAXONOMY_VERSION, CSRD_TAXONOMY_REFERENCE,
        CSRD_CATEGORIES, ESRS_SUBCATEGORIES, MATERIALITY_SCALE,
        MARKET_SURPRISE_LEVELS, TIME_HORIZONS,
    )
except ImportError:
    print(
        "❌ csrd_taxonomy.py introuvable.\n"
        "   Placez csrd_taxonomy.py dans le même dossier que ce script."
    )
    sys.exit(1)

# ═════════════════════════════════════════════════════════════════════════════
# CONFIGURATION
# ═════════════════════════════════════════════════════════════════════════════

N_GOLD = 500
N_IAA  = 200
RANDOM_SEED = 42                  # Cohérent avec Phase 2 — reproductibilité stricte
MIN_PER_STRATUM_GOLD = 15         # 500 / 16 strates ≈ 31 en moyenne
MIN_PER_STRATUM_IAA  = 6          # 200 / 16 strates ≈ 12.5 en moyenne

YEAR_BUCKETS_ORDER = ["2010-2014", "2015-2019", "2020-2022", "2023-2026"]
CSRD_QUARTILE_ORDER = ["q0_low", "q1_med_low", "q2_med_high", "q3_high"]

# ═════════════════════════════════════════════════════════════════════════════
# CHEMINS
# ═════════════════════════════════════════════════════════════════════════════

ROOT_DIR  = Path(__file__).parent
DATA_DIR  = ROOT_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
RUN_TS   = datetime.now().strftime("%Y%m%d_%H%M%S")

OUT_XLSX  = DATA_DIR / f"frafin_gold_500_{RUN_TS}.xlsx"
OUT_REPORT = DATA_DIR / f"gold_sampling_report_{RUN_TS}.json"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger(__name__)

# ═════════════════════════════════════════════════════════════════════════════
# CHARGEMENT
# ═════════════════════════════════════════════════════════════════════════════

def load_pool(filepath: str) -> pd.DataFrame:
    path = Path(filepath)
    if not path.exists() or filepath == "auto":
        candidates = sorted(DATA_DIR.glob("frafin_sample_rigorous_*.parquet"), reverse=True)
        if not candidates:
            candidates = sorted(DATA_DIR.glob("frafin_sample_rigorous_*.csv"), reverse=True)
        if not candidates:
            logger.error(
                "Aucun fichier frafin_sample_rigorous_*.parquet trouvé dans data/. "
                "Lancez d'abord build_annotation_sample_rigorous.py (Phase 2)."
            )
            sys.exit(1)
        path = candidates[0]
        logger.info(f"Auto-détection : {path.name}")

    df = pd.read_parquet(path) if path.suffix == ".parquet" else \
         pd.read_csv(path, dtype=str, low_memory=False)

    required = {"paragraph_id", "content", "year_bucket", "csrd_quartile"}
    missing  = required - set(df.columns)
    if missing:
        logger.error(
            f"Colonnes requises manquantes : {missing}. "
            f"Ce script attend la sortie de build_annotation_sample_rigorous.py."
        )
        sys.exit(1)

    logger.info(f"Pool chargé : {len(df):,} paragraphes  ({path.name})")
    return df

# ═════════════════════════════════════════════════════════════════════════════
# ALLOCATION DE QUOTAS PAR STRATE (réutilisable GOLD et IAA)
# ═════════════════════════════════════════════════════════════════════════════

def allocate_quota(
    strata_sizes: pd.Series, target_total: int, min_per_stratum: int
) -> dict[str, int]:
    """
    Allocation proportionnelle avec plancher minimum, ajustée pour atteindre
    exactement target_total (compense les arrondis et les planchers).
    """
    total_pool = int(strata_sizes.sum())
    quotas: dict[str, int] = {}
    for stratum, size in strata_sizes.items():
        q = max(min_per_stratum, int(round(target_total * size / total_pool)))
        quotas[stratum] = min(q, int(size))

    total_alloc = sum(quotas.values())

    # Trop alloué → réduire les strates au-dessus du plancher (les plus grandes d'abord)
    if total_alloc > target_total:
        excess = total_alloc - target_total
        reducible = sorted(
            [s for s in quotas if quotas[s] > min_per_stratum],
            key=lambda s: quotas[s], reverse=True,
        )
        i = 0
        while excess > 0 and reducible:
            s = reducible[i % len(reducible)]
            if quotas[s] > min_per_stratum:
                quotas[s] -= 1
                excess -= 1
            i += 1
            if i > 200_000:
                break

    # Pas assez alloué → compléter sur les strates ayant encore de la marge
    elif total_alloc < target_total:
        deficit = target_total - total_alloc
        capacity = {s: int(strata_sizes[s]) - quotas[s] for s in quotas}
        capacity = {s: c for s, c in capacity.items() if c > 0}
        order = sorted(capacity, key=lambda s: capacity[s], reverse=True)
        i = 0
        while deficit > 0 and order:
            s = order[i % len(order)]
            if capacity[s] > 0:
                quotas[s] += 1
                capacity[s] -= 1
                deficit -= 1
            i += 1
            if i > 200_000:
                break

    return quotas

# ═════════════════════════════════════════════════════════════════════════════
# DIVERSIFICATION ROUND-ROBIN PAR TYPE DE DOCUMENT
# ═════════════════════════════════════════════════════════════════════════════

def diversify_round_robin(
    pool: pd.DataFrame, n: int, group_col: str, seed: int
) -> pd.DataFrame:
    """
    Sélectionne n lignes de `pool` en alternant entre les groupes de group_col,
    pour maximiser la diversité de types de documents au sein d'une strate,
    plutôt que de piocher au hasard pur (qui pourrait sur-représenter le
    type de document le plus fréquent dans la strate).
    """
    if len(pool) <= n or group_col not in pool.columns:
        return pool.sample(frac=1, random_state=seed).head(n).reset_index(drop=True) \
               if len(pool) > n else pool.copy()

    rng = np.random.RandomState(seed)
    groups: dict[str, deque] = {}
    for g, sub in pool.groupby(group_col):
        shuffled = sub.sample(frac=1, random_state=seed).reset_index(drop=True)
        groups[g] = deque(shuffled.to_dict("records"))

    order = list(groups.keys())
    rng.shuffle(order)

    selected: list[dict] = []
    while len(selected) < n:
        progressed = False
        for g in order:
            if groups[g]:
                selected.append(groups[g].popleft())
                progressed = True
                if len(selected) >= n:
                    break
        if not progressed:
            break

    return pd.DataFrame(selected).reset_index(drop=True)

# ═════════════════════════════════════════════════════════════════════════════
# SÉLECTION DES 500 GOLD
# ═════════════════════════════════════════════════════════════════════════════

def select_gold(
    df: pd.DataFrame, n_gold: int, min_per_stratum: int, seed: int
) -> tuple[pd.DataFrame, dict]:
    logger.info(f"\n{'═'*65}")
    logger.info(f"  SÉLECTION DES {n_gold} GOLD — Stratification "
                f"year_bucket × csrd_quartile")
    logger.info(f"{'═'*65}")

    df = df.copy()
    df["gold_stratum"] = (
        df["year_bucket"].astype(str) + "__" + df["csrd_quartile"].astype(str)
    )
    strata_sizes = df["gold_stratum"].value_counts()
    logger.info(f"  Strates actives : {len(strata_sizes)} (attendu 16 max)")

    quotas = allocate_quota(strata_sizes, n_gold, min_per_stratum)

    parts = []
    for stratum, quota in quotas.items():
        pool = df[df["gold_stratum"] == stratum]
        picked = diversify_round_robin(pool, quota, "doc_group", seed)
        parts.append(picked)

    gold_df = pd.concat(parts, ignore_index=True)
    gold_df = gold_df.sample(frac=1, random_state=seed).reset_index(drop=True)

    logger.info(f"\n  Total GOLD sélectionné : {len(gold_df):,}")
    logger.info(f"\n  Distribution year_bucket × csrd_quartile :")
    pivot = pd.crosstab(gold_df["year_bucket"], gold_df["csrd_quartile"])
    pivot = pivot.reindex(index=YEAR_BUCKETS_ORDER, columns=CSRD_QUARTILE_ORDER, fill_value=0)
    logger.info(f"\n{pivot.to_string()}")

    if "doc_group" in gold_df.columns:
        logger.info(f"\n  Types de documents couverts : {gold_df['doc_group'].nunique()}")
        for grp, cnt in gold_df["doc_group"].value_counts().head(10).items():
            logger.info(f"    {str(grp):<40} {cnt:>4}")

    return gold_df, quotas

# ═════════════════════════════════════════════════════════════════════════════
# SÉLECTION DU SOUS-ENSEMBLE IAA (200 / 500)
# ═════════════════════════════════════════════════════════════════════════════

def select_iaa_subset(
    gold_df: pd.DataFrame, n_iaa: int, min_per_stratum: int, seed: int
) -> tuple[pd.DataFrame, dict]:
    logger.info(f"\n{'═'*65}")
    logger.info(f"  SÉLECTION DU SOUS-ENSEMBLE IAA ({n_iaa} / {len(gold_df)})")
    logger.info(f"{'═'*65}")

    strata_sizes = gold_df["gold_stratum"].value_counts()
    quotas = allocate_quota(strata_sizes, n_iaa, min_per_stratum)

    iaa_ids: set[str] = set()
    for stratum, quota in quotas.items():
        pool = gold_df[gold_df["gold_stratum"] == stratum]
        n_pick = min(quota, len(pool))
        picked = pool.sample(n=n_pick, random_state=seed)
        iaa_ids.update(picked["paragraph_id"])

    gold_df = gold_df.copy()
    gold_df["iaa_subset"] = gold_df["paragraph_id"].isin(iaa_ids)

    n_marked = int(gold_df["iaa_subset"].sum())
    logger.info(f"  Paragraphes marqués IAA : {n_marked} / {len(gold_df)}")
    logger.info(f"  (seront ré-annotés indépendamment par Mistral pour le Kappa de Cohen)")

    return gold_df, quotas

# ═════════════════════════════════════════════════════════════════════════════
# STYLES EXCEL PARTAGÉS
# ═════════════════════════════════════════════════════════════════════════════

FONT_NAME = "Arial"

STYLE_HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
FILL_HEADER  = PatternFill("solid", fgColor="1F4E78")
FILL_READONLY = PatternFill("solid", fgColor="F2F2F2")
FILL_EDITABLE = PatternFill("solid", fgColor="FFF9DB")
FILL_INFO     = PatternFill("solid", fgColor="DCE6F1")
FONT_NORMAL   = Font(name=FONT_NAME, size=10)
FONT_ITALIC   = Font(name=FONT_NAME, size=9, italic=True, color="666666")
THIN_BORDER   = Border(*([Side(style="thin", color="D9D9D9")] * 4))
ALIGN_WRAP    = Alignment(wrap_text=True, vertical="top", horizontal="left")
ALIGN_CENTER  = Alignment(horizontal="center", vertical="center")


def style_header_row(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        c = ws.cell(row=row, column=col)
        c.font = STYLE_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Instructions"
# ═════════════════════════════════════════════════════════════════════════════

def build_instructions_sheet(wb: Workbook, n_gold: int, n_iaa: int) -> None:
    ws = wb.create_sheet("Instructions")
    ws.column_dimensions["A"].width = 100
    ws.sheet_view.showGridLines = False

    lines = [
        ("FraFin-Reasoning — Annotation Gold Standard (Hold-out Test Set)", True, 14),
        ("", False, 10),
        (f"Ce classeur contient {n_gold} paragraphes issus du corpus réglementaire AMF "
         f"(CAC 40, 2010–2026), sélectionnés par échantillonnage stratifié proportionnel "
         f"(year_bucket × csrd_quartile) — PAS par sélection des cas les plus intéressants.",
         False, 10),
        ("", False, 10),
        ("CE QUE VOUS DEVEZ FAIRE", True, 12),
        ("→ Onglet 'Annotation' : remplissez les colonnes en fond jaune (H à Q) "
         "pour chaque ligne, en lisant l'extrait en colonne G.",
         False, 10),
        ("→ Les colonnes en fond gris sont en lecture seule (métadonnées).", False, 10),
        ("→ La colonne en fond bleu (IAA) est informative : elle indique quelles lignes "
         "seront aussi annotées indépendamment par Mistral pour le calcul du Kappa de Cohen.",
         False, 10),
        ("", False, 10),
        ("COLONNES À REMPLIR", True, 12),
        ("• csrd_category — catégorie ESRS principale, ou 'none' si aucune pertinence "
         "CSRD/ESG. Voir l'onglet 'Taxonomie_CSRD'.", False, 10),
        ("• esrs_subcategory — se met à jour automatiquement selon la catégorie choisie "
         "(menu déroulant dépendant). Si le menu ne se filtre pas automatiquement selon "
         "votre version d'Excel, consultez l'onglet Taxonomie_CSRD et saisissez le code "
         "manuellement.", False, 10),
        ("• materialite_score — 0 si csrd_category='none', sinon 1 (négligeable) à "
         "5 (critique). Voir l'onglet 'Echelle_Materialite'.", False, 10),
        ("• materialite_financiere / materialite_impact — double matérialité CSRD. "
         "'Oui'/'Non' indépendamment l'un de l'autre. Les deux doivent être 'Non' si "
         "csrd_category='none'.", False, 10),
        ("• market_surprise — niveau de surprise par rapport au consensus de marché : "
         "aucune / faible / partielle / forte.", False, 10),
        ("• horizon_temporel — immediat / court_terme / moyen_terme / long_terme. "
         "Laisser vide si csrd_category='none'.", False, 10),
        ("• chain_of_thought — votre raisonnement en 2-3 phrases : (1) nature de "
         "l'information, (2) analyse de matérialité, (3) évaluation de la surprise.",
         False, 10),
        ("• confiance_annotation — votre confiance dans l'annotation, de 0.0 à 1.0.",
         False, 10),
        ("• notes_annotateur (optionnel) — cas ambigus, hésitations, à discuter pour "
         "affiner les guidelines.", False, 10),
        ("", False, 10),
        ("RÈGLE DE COHÉRENCE 'NONE'", True, 12),
        ("Si csrd_category = 'none' : esrs_subcategory vide, materialite_score = 0, "
         "materialite_financiere = Non, materialite_impact = Non, horizon_temporel vide. "
         "Ne forcez jamais une catégorie CSRD sur un paragraphe purement financier ou "
         "administratif sans dimension de durabilité — c'est la classe la plus "
         "importante du dataset.", False, 10),
        ("", False, 10),
        ("PROTOCOLE D'ANNOTATION À L'AVEUGLE", True, 12),
        ("Les variables qui ont servi à sélectionner cet échantillon (quartile CSRD, "
         "score de mots-clés) ne sont PAS affichées dans cette feuille, pour éviter "
         "tout biais d'ancrage dans votre jugement. Annotez chaque paragraphe selon "
         "son seul contenu.", False, 10),
        ("", False, 10),
        (f"Sous-ensemble IAA : {n_iaa} des {n_gold} paragraphes (colonne 'iaa_subset') "
         f"seront également annotés indépendamment par Mistral Large pour mesurer "
         f"l'accord inter-annotateurs (Kappa de Cohen, cible κ > 0.75).", False, 10),
    ]

    r = 1
    for text, bold, size in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT_NAME, size=size, bold=bold,
                      color="1F4E78" if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.row_dimensions[1].height = 24

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Taxonomie_CSRD"
# ═════════════════════════════════════════════════════════════════════════════

def build_taxonomy_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Taxonomie_CSRD")
    headers = ["Code", "Pilier", "Libellé (FR)", "Description", "Sous-catégories (codes)"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
    style_header_row(ws, 1, len(headers))

    widths = [10, 14, 38, 70, 60]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    r = 2
    for code, info in CSRD_CATEGORIES.items():
        subs = ", ".join(s["code"] for s in ESRS_SUBCATEGORIES.get(code, []))
        row_vals = [code, info["pilier"], info["label_fr"], info["description"], subs]
        for col, val in enumerate(row_vals, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = FONT_NORMAL
            c.alignment = ALIGN_WRAP
            c.border = THIN_BORDER
        ws.row_dimensions[r].height = 60
        r += 1

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    # Note de version en bas
    note = ws.cell(
        row=r + 1, column=1,
        value=f"Taxonomie v{CSRD_TAXONOMY_VERSION} — {CSRD_TAXONOMY_REFERENCE}",
    )
    note.font = FONT_ITALIC

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Echelle_Materialite"
# ═════════════════════════════════════════════════════════════════════════════

def build_materiality_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Echelle_Materialite")
    headers = ["Score", "Libellé", "Description"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
    style_header_row(ws, 1, len(headers))

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 80

    r = 2
    for score, info in MATERIALITY_SCALE.items():
        ws.cell(row=r, column=1, value=score).alignment = ALIGN_CENTER
        ws.cell(row=r, column=2, value=info["label_fr"])
        c3 = ws.cell(row=r, column=3, value=info["description"])
        c3.alignment = ALIGN_WRAP
        for col in range(1, 4):
            ws.cell(row=r, column=col).font = FONT_NORMAL
            ws.cell(row=r, column=col).border = THIN_BORDER
        ws.row_dimensions[r].height = 40
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Surprise de marché").font = Font(
        name=FONT_NAME, bold=True, size=11
    )
    r += 1
    for code, info in MARKET_SURPRISE_LEVELS.items():
        ws.cell(row=r, column=1, value=code).font = FONT_NORMAL
        c = ws.cell(row=r, column=2, value=info["description"])
        c.alignment = ALIGN_WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Horizon temporel").font = Font(
        name=FONT_NAME, bold=True, size=11
    )
    r += 1
    for code, info in TIME_HORIZONS.items():
        ws.cell(row=r, column=1, value=code).font = FONT_NORMAL
        c = ws.cell(row=r, column=2, value=info["description"])
        c.alignment = ALIGN_WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 30
        r += 1

    ws.sheet_view.showGridLines = False

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Lists" (masquée) — sources des menus déroulants en cascade
# ═════════════════════════════════════════════════════════════════════════════

def build_lists_sheet(wb: Workbook) -> dict[str, str]:
    """
    Écrit les sous-catégories de chaque catégorie ESRS dans une colonne dédiée,
    crée une plage nommée SUB_<code> par catégorie, pour alimenter le menu
    déroulant en cascade esrs_subcategory (dépendant de csrd_category).
    Retourne {category_code: defined_name}.
    """
    ws = wb.create_sheet("Lists")
    ws.sheet_state = "hidden"

    defined_names: dict[str, str] = {}

    for col_idx, (code, subs) in enumerate(ESRS_SUBCATEGORIES.items(), start=1):
        col_letter = get_column_letter(col_idx)
        ws.cell(row=1, column=col_idx, value=code)   # en-tête = code catégorie
        for i, s in enumerate(subs, start=2):
            ws.cell(row=i, column=col_idx, value=s["code"])

        last_row = 1 + len(subs)
        name = f"SUB_{code}"
        ref  = f"Lists!${col_letter}$2:${col_letter}${last_row}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))
        defined_names[code] = name

    # Liste plate de toutes les catégories (pour le menu déroulant csrd_category)
    cat_col = len(ESRS_SUBCATEGORIES) + 2
    cat_letter = get_column_letter(cat_col)
    ws.cell(row=1, column=cat_col, value="ALL_CATEGORIES")
    for i, code in enumerate(CSRD_CATEGORIES.keys(), start=2):
        ws.cell(row=i, column=cat_col, value=code)
    wb.defined_names.add(DefinedName(
        "ALL_CATEGORIES",
        attr_text=f"Lists!${cat_letter}$2:${cat_letter}${1+len(CSRD_CATEGORIES)}",
    ))

    return defined_names

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Annotation" (principale)
# ═════════════════════════════════════════════════════════════════════════════

# Colonnes : lettre → (clé_df_ou_None, libellé, largeur, type)
# type : "readonly" | "editable" | "info"
COLUMNS: list[tuple[str, str | None, str, int, str]] = [
    ("A", "_idx",           "#",                    5,  "readonly"),
    ("B", "paragraph_id",   "paragraph_id",          14, "readonly"),
    ("C", "author",         "Société",               22, "readonly"),
    ("D", "_year",          "Année",                  7, "readonly"),
    ("E", "doc_type",       "Type de document",      30, "readonly"),
    ("F", "source_url",     "Source (lien)",         12, "readonly"),
    ("G", "content",        "Extrait à annoter",     70, "readonly"),
    ("H", None,             "csrd_category",         16, "editable"),
    ("I", None,             "esrs_subcategory",      16, "editable"),
    ("J", None,             "materialite_score",     12, "editable"),
    ("K", None,             "materialite_financiere",14, "editable"),
    ("L", None,             "materialite_impact",    14, "editable"),
    ("M", None,             "market_surprise",       14, "editable"),
    ("N", None,             "horizon_temporel",      14, "editable"),
    ("O", None,             "chain_of_thought",      55, "editable"),
    ("P", None,             "confiance_annotation",  12, "editable"),
    ("Q", None,             "notes_annotateur",      30, "editable"),
    ("R", "iaa_subset",     "IAA ?",                 10, "info"),
]


def build_annotation_sheet(wb: Workbook, gold_df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Annotation")
    ws.sheet_view.showGridLines = False

    # ── En-têtes ───────────────────────────────────────────────────────────
    for letter, _, label, width, _ in COLUMNS:
        col = ws[f"{letter}1"]
        col.value = label
        ws.column_dimensions[letter].width = width
    style_header_row(ws, 1, len(COLUMNS))
    ws.row_dimensions[1].height = 30

    # ── Données ────────────────────────────────────────────────────────────
    n_rows = len(gold_df)
    for i, (_, row) in enumerate(gold_df.iterrows(), start=2):
        for letter, key, _, _, kind in COLUMNS:
            if key == "_idx":
                val = i - 1
            elif key == "iaa_subset":
                val = "Oui" if bool(row.get("iaa_subset", False)) else "Non"
            elif key == "_year":
                y = row.get("_year", 0)
                val = int(y) if pd.notna(y) and y else None
            elif key is not None:
                val = row.get(key)
                if pd.isna(val):
                    val = None
            else:
                val = None   # colonnes éditables, vides au départ

            cell = ws[f"{letter}{i}"]
            cell.value = val
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER

            if kind == "readonly":
                cell.fill = FILL_READONLY
            elif kind == "editable":
                cell.fill = FILL_EDITABLE
            elif kind == "info":
                cell.fill = FILL_INFO

            if letter in ("G", "O"):
                cell.alignment = ALIGN_WRAP
            elif letter in ("A", "D", "J", "P", "R"):
                cell.alignment = ALIGN_CENTER

        # Hyperlien sur la source
        url = row.get("source_url")
        if isinstance(url, str) and url.startswith("http"):
            fc = ws[f"F{i}"]
            fc.value = "Voir document"
            fc.hyperlink = url
            fc.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

        ws.row_dimensions[i].height = 70

    last_row = n_rows + 1

    # ── Validations de données ────────────────────────────────────────────
    def add_list_dv(col_letter: str, formula: str, allow_blank: bool = False,
                     prompt: str = "", title: str = "") -> None:
        dv = DataValidation(
            type="list", formula1=formula, allow_blank=allow_blank,
            showDropDown=False, showErrorMessage=True,
            errorTitle="Valeur invalide", error="Choisissez une valeur de la liste.",
            showInputMessage=bool(prompt), promptTitle=title or None, prompt=prompt or None,
        )
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{last_row}")

    cat_list = ",".join(CSRD_CATEGORIES.keys())
    add_list_dv("H", f'"{cat_list}"',
                title="Catégorie CSRD",
                prompt="Choisir 'none' si aucune pertinence CSRD/ESG.")

    # Sous-catégorie en cascade : dépend de la cellule H de la même ligne
    dv_sub = DataValidation(
        type="list", formula1='INDIRECT("SUB_"&H2)', allow_blank=True,
        showErrorMessage=True, errorTitle="Valeur invalide",
        error="Choisissez d'abord une catégorie en colonne H, puis une sous-catégorie cohérente.",
        showInputMessage=True, promptTitle="Sous-catégorie ESRS",
        prompt="La liste dépend de la catégorie choisie en colonne H.",
    )
    ws.add_data_validation(dv_sub)
    dv_sub.add(f"I2:I{last_row}")

    add_list_dv("J", '"0,1,2,3,4,5"', title="Matérialité",
                prompt="0 si csrd_category='none', sinon 1 à 5.")
    add_list_dv("K", '"Oui,Non"')
    add_list_dv("L", '"Oui,Non"')

    surprise_list = ",".join(MARKET_SURPRISE_LEVELS.keys())
    add_list_dv("M", f'"{surprise_list}"')

    horizon_list = ",".join(TIME_HORIZONS.keys())
    add_list_dv("N", f'"{horizon_list}"', allow_blank=True,
                prompt="Laisser vide si csrd_category='none'.")

    dv_conf = DataValidation(
        type="decimal", operator="between", formula1=0, formula2=1,
        allow_blank=True, showErrorMessage=True,
        errorTitle="Valeur invalide", error="Confiance entre 0.0 et 1.0.",
    )
    ws.add_data_validation(dv_conf)
    dv_conf.add(f"P2:P{last_row}")

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:R{last_row}"

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE INTERNE "Metadonnees_Stratification" (masquée — pour le chercheur)
# ═════════════════════════════════════════════════════════════════════════════

def build_internal_metadata_sheet(wb: Workbook, gold_df: pd.DataFrame) -> None:
    """
    Conserve les variables ayant servi à la sélection (csrd_quartile, score,
    stratum) hors de la vue de l'annotateur, pour préserver l'aveuglement.
    Réservée à vos propres besoins de traçabilité méthodologique.
    """
    ws = wb.create_sheet("Metadonnees_Stratification")
    ws.sheet_state = "hidden"

    cols = ["paragraph_id", "gold_stratum", "year_bucket", "csrd_quartile",
            "csrd_score_raw", "stratum_id", "doc_group", "iaa_subset"]
    cols = [c for c in cols if c in gold_df.columns or c == "iaa_subset"]

    for col_idx, name in enumerate(cols, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    for r, (_, row) in enumerate(gold_df.iterrows(), start=2):
        for col_idx, name in enumerate(cols, start=1):
            val = row.get(name)
            if isinstance(val, (np.bool_,)):
                val = bool(val)
            if pd.isna(val) if not isinstance(val, bool) else False:
                val = None
            ws.cell(row=r, column=col_idx, value=val)

# ═════════════════════════════════════════════════════════════════════════════
# FEUILLE "Resume_Strates"
# ═════════════════════════════════════════════════════════════════════════════

def build_summary_sheet(wb: Workbook, gold_df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Resume_Strates")
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Distribution des 500 GOLD — year_bucket × csrd_quartile"
    ws["A1"].font = Font(name=FONT_NAME, bold=True, size=12)

    pivot = pd.crosstab(gold_df["year_bucket"], gold_df["csrd_quartile"])
    pivot = pivot.reindex(index=YEAR_BUCKETS_ORDER, columns=CSRD_QUARTILE_ORDER, fill_value=0)

    ws.cell(row=3, column=1, value="year_bucket")
    for j, col in enumerate(pivot.columns, start=2):
        ws.cell(row=3, column=j, value=col)
    style_header_row(ws, 3, len(pivot.columns) + 1)

    for i, (idx, vals) in enumerate(pivot.iterrows(), start=4):
        ws.cell(row=i, column=1, value=idx).font = Font(name=FONT_NAME, bold=True)
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=i, column=j, value=int(v))
            c.font = FONT_NORMAL
            c.alignment = ALIGN_CENTER

    for col in range(1, len(pivot.columns) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 16

    if "doc_group" in gold_df.columns:
        r0 = len(pivot) + 6
        ws.cell(row=r0, column=1,
                value="Distribution par type de document").font = Font(
                    name=FONT_NAME, bold=True, size=12)
        r0 += 1
        ws.cell(row=r0, column=1, value="doc_group")
        ws.cell(row=r0, column=2, value="n")
        style_header_row(ws, r0, 2)
        r0 += 1
        for grp, cnt in gold_df["doc_group"].value_counts().items():
            ws.cell(row=r0, column=1, value=str(grp)).font = FONT_NORMAL
            ws.cell(row=r0, column=2, value=int(cnt)).font = FONT_NORMAL
            r0 += 1

# ═════════════════════════════════════════════════════════════════════════════
# RAPPORT JSON
# ═════════════════════════════════════════════════════════════════════════════

def save_report(
    pool_size: int, gold_df: pd.DataFrame,
    gold_quotas: dict, iaa_quotas: dict,
) -> None:
    pivot = pd.crosstab(gold_df["year_bucket"], gold_df["csrd_quartile"])
    pivot = pivot.reindex(index=YEAR_BUCKETS_ORDER, columns=CSRD_QUARTILE_ORDER, fill_value=0)

    report = {
        "run_timestamp": RUN_TS,
        "random_seed":   RANDOM_SEED,
        "taxonomy_version": CSRD_TAXONOMY_VERSION,
        "pool_size":     pool_size,
        "n_gold":        len(gold_df),
        "n_iaa":         int(gold_df["iaa_subset"].sum()),
        "min_per_stratum_gold": MIN_PER_STRATUM_GOLD,
        "min_per_stratum_iaa":  MIN_PER_STRATUM_IAA,
        "stratification_axes":  ["year_bucket", "csrd_quartile"],
        "diversification":      "round_robin by doc_group within each stratum",
        "blind_annotation": (
            "csrd_quartile, csrd_score_raw and stratum_id are withheld from the "
            "annotation sheet shown to the human annotator to prevent anchoring "
            "bias; preserved in a hidden internal sheet for traceability."
        ),
        "gold_quotas_by_stratum": gold_quotas,
        "iaa_quotas_by_stratum":  iaa_quotas,
        "distribution_year_x_quartile": pivot.to_dict(),
        "doc_group_distribution": (
            gold_df["doc_group"].value_counts().to_dict()
            if "doc_group" in gold_df.columns else {}
        ),
        "output_files": {
            "excel":  str(OUT_XLSX),
            "report": str(OUT_REPORT),
        },
        "arxiv_text": (
            f"The Gold Standard hold-out test set ({len(gold_df)} paragraphs) was "
            f"selected via proportional stratified sampling on two axes (temporal "
            f"bucket × CSRD density quartile, 16 strata) with a minimum guarantee "
            f"of {MIN_PER_STRATUM_GOLD} paragraphs per stratum, and document-type "
            f"diversification via round-robin selection within each stratum. "
            f"This avoids the selection bias of choosing only the most semantically "
            f"rich examples, which would overestimate downstream model performance "
            f"on this test set. The stratification variables (CSRD density quartile, "
            f"raw score) were withheld from the human annotator during labeling to "
            f"prevent anchoring bias. A {int(gold_df['iaa_subset'].sum())}-paragraph "
            f"subset, itself stratified across the same two axes, was independently "
            f"double-annotated by Mistral Large to compute inter-annotator agreement "
            f"(Cohen's Kappa)."
        ),
    }

    with open(OUT_REPORT, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False, default=str)
    logger.info(f"  ✅ Rapport : {OUT_REPORT}")

# ═════════════════════════════════════════════════════════════════════════════
# CLI & MAIN
# ═════════════════════════════════════════════════════════════════════════════

def parse_args():
    p = argparse.ArgumentParser(
        description="FraFin-Reasoning — Export Excel des 500 Gold pour annotation humaine"
    )
    p.add_argument("--input", default="auto")
    p.add_argument("--n-gold", type=int, default=N_GOLD)
    p.add_argument("--n-iaa",  type=int, default=N_IAA)
    p.add_argument("--seed",   type=int, default=RANDOM_SEED)
    p.add_argument("--min-per-stratum-gold", type=int, default=MIN_PER_STRATUM_GOLD)
    p.add_argument("--min-per-stratum-iaa",  type=int, default=MIN_PER_STRATUM_IAA)
    return p.parse_args()


def main():
    global N_GOLD, N_IAA, RANDOM_SEED, MIN_PER_STRATUM_GOLD, MIN_PER_STRATUM_IAA

    args = parse_args()
    N_GOLD  = args.n_gold
    N_IAA   = args.n_iaa
    RANDOM_SEED = args.seed
    MIN_PER_STRATUM_GOLD = args.min_per_stratum_gold
    MIN_PER_STRATUM_IAA  = args.min_per_stratum_iaa

    logger.info("═" * 65)
    logger.info("  FraFin-Reasoning — Export Gold Standard (Phase 3a)")
    logger.info(f"  GOLD : {N_GOLD}  |  IAA : {N_IAA}  |  Seed : {RANDOM_SEED}")
    logger.info("═" * 65)

    df = load_pool(args.input)
    pool_size = len(df)

    gold_df, gold_quotas = select_gold(df, N_GOLD, MIN_PER_STRATUM_GOLD, RANDOM_SEED)
    gold_df, iaa_quotas  = select_iaa_subset(gold_df, N_IAA, MIN_PER_STRATUM_IAA, RANDOM_SEED)

    logger.info(f"\n{'═'*65}")
    logger.info(f"  CONSTRUCTION DU CLASSEUR EXCEL")
    logger.info(f"{'═'*65}")

    wb = Workbook()
    wb.remove(wb.active)   # supprime la feuille vide par défaut
    wb.properties.title   = "FraFin-Reasoning — Gold Standard Annotation"
    wb.properties.creator = "FraFin-Reasoning Project (ENSAE)"

    build_instructions_sheet(wb, N_GOLD, N_IAA)
    build_taxonomy_sheet(wb)
    build_materiality_sheet(wb)
    build_lists_sheet(wb)
    build_annotation_sheet(wb, gold_df)
    build_internal_metadata_sheet(wb, gold_df)
    build_summary_sheet(wb, gold_df)

    # Onglet actif à l'ouverture = Annotation
    wb.active = wb.sheetnames.index("Annotation")

    wb.save(OUT_XLSX)
    logger.info(f"  ✅ Excel : {OUT_XLSX}")

    save_report(pool_size, gold_df, gold_quotas, iaa_quotas)

    logger.info(f"\n{'═'*65}")
    logger.info(f"  TERMINÉ")
    logger.info(f"{'═'*65}")
    logger.info(f"  Ouvrez {OUT_XLSX.name} et complétez l'onglet 'Annotation'.")
    logger.info(f"  ⚙  Prochaine étape : annotation_mistral_mass.py")
    logger.info("═" * 65)


if __name__ == "__main__":
    main()
